# -*- coding: utf-8 -*-
"""原価表 Excel（BFR/PLD/BFQ 原価ブック）の解析

ブック構成（3ブック共通。docs/原価検証システム_要件整理と実装方針_20260907.md 1.1 参照）:
  - 単価マスタシート: C1 が「鋼材単価表」。上段=鋼材単価表（鋼種×板厚→円/kg）、
    下段=部品リスト（B=原価表上名称, C/D=鋼種・板厚（鋼板の別名のとき）, E=旧単価, F=新単価, G=根拠）
  - 型式シート: 「内訳」ヘッダー行の下に、部位ごとの明細行。列は見出し文字で判定する。

このモジュールは DB に依存しない。parse_workbook() が返す dict を
  - API の取込プレビュー（build_context → costing_engine で Excel 値と突合）
  - API の取込確定（DB へ書込み）
の両方で使う。
"""
from __future__ import annotations
import re
import hashlib
import datetime as dt
from collections import OrderedDict

try:
    import openpyxl
except ImportError:  # requirements.txt に openpyxl を追加済み。無い環境では取込だけ使えない
    openpyxl = None

# 単価の時点（マスタ列見出し「2025年9月時点単価」→ 2025-09-01）
_DATE_RE = re.compile(r"(\d{4})\s*年\s*(\d{1,2})\s*月")
_SHEET_REF_RE = re.compile(r"'?([^'!=+*/(),]+?)'?!\$?[A-Z]{1,3}\$?\d+")
_SUM_RE = re.compile(r"SUM\(\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)\)", re.IGNORECASE)
_CELL_RE = re.compile(r"(?<![A-Z'!])\$?([A-Z]{1,3})\$?(\d+)(?![\d(])")
_TECHS_RE = re.compile(r"^\d{6,12}$")
_PAINT_RE = re.compile(r"ﾏﾘﾝ|マリン|ｼﾝﾅｰ|シンナー|ﾌﾟﾗｲﾏｰ|プライマー|ﾘﾙｶ|リルカ|塗料|ﾍﾟｲﾝﾄ")

CATEGORY_STEEL = "鋼板"
CATEGORY_AREA = "面積材"
CATEGORY_SHAPE = "形鋼"
CATEGORY_PIPE = "パイプ"
CATEGORY_PURCHASE = "購入品"
CATEGORY_SUBCON = "外注加工"
CATEGORY_PAINT = "塗料"
CATEGORIES = [CATEGORY_STEEL, CATEGORY_AREA, CATEGORY_SHAPE, CATEGORY_PIPE,
              CATEGORY_PURCHASE, CATEGORY_SUBCON, CATEGORY_PAINT]

# 種別 → 標準原価の「鋼材」側に集計するか
STEEL_CATEGORIES = {CATEGORY_STEEL, CATEGORY_AREA, CATEGORY_SHAPE, CATEGORY_PIPE}


# ------------------------------------------------------------
# ユーティリティ
# ------------------------------------------------------------
def norm_label(v) -> str:
    """原価表上名称の正規化。数値（板厚 3.2 / 6）は Excel の表示どおりの文字列にする。"""
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        f = float(v)
        return str(int(f)) if f.is_integer() else repr(f)
    s = str(v).replace("　", " ").strip()
    return re.sub(r"\s+", " ", s)


def temp_code(name: str) -> str:
    """TECHS 未採番の資材に付ける仮コード（名称から決定的に生成。再取込で同じコードになる）"""
    return "T-" + hashlib.sha1(name.encode("utf-8")).hexdigest()[:8].upper()


def _is_num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _parse_date_header(text) -> dt.date | None:
    if not isinstance(text, str):
        return None
    m = _DATE_RE.search(text)
    if not m:
        return None
    return dt.date(int(m.group(1)), int(m.group(2)), 1)


def _price_unit_from_basis(basis: str, default: str) -> str:
    """「(円/kg)×21.2kg/6.0m=(円/m)」のような根拠文字列から単価の単位を読む（末尾の =(円/x) を優先）"""
    if not isinstance(basis, str):
        return default
    tail = re.findall(r"円/(kg|㎡|m2|m|本|枚|個|缶|ﾘﾝｸ|リンク)", basis)
    if not tail:
        return default
    u = tail[-1]
    return {"㎡": "m2", "ﾘﾝｸ": "リンク"}.get(u, u)


def classify_material(name: str, basis: str = "") -> tuple[str, str]:
    """名称と根拠から (種別, 単価単位) を推定する"""
    n = name
    if "/外注/" in n or n.startswith("外注/") or n.startswith("外注-"):
        return CATEGORY_SUBCON, "個"
    if _PAINT_RE.search(n):
        return CATEGORY_PAINT, "缶"
    if re.match(r"^(FB=|ｱﾝｸﾞﾙ|アングル|STKR|ﾁｬﾝﾈﾙ|チャンネル|SS黒丸|SS磨丸|SUS304 引抜|SKロット|ﾕﾆｸﾛ寸切)", n):
        return CATEGORY_SHAPE, _price_unit_from_basis(basis, "m")
    if re.match(r"^(SGP|一般構造用|継目無鋼管|STKM|ﾊﾟｲﾌﾟ|パイプ)", n):
        return CATEGORY_PIPE, _price_unit_from_basis(basis, "m")
    if re.match(r"^(ﾊﾟﾝﾁﾝｸﾞ|パンチング|ｴｷｽﾊﾟﾝﾄﾞ|エキスパンド|ﾎﾞﾝﾃﾞﾊﾟﾝﾁﾝｸﾞ)", n):
        return CATEGORY_AREA, _price_unit_from_basis(basis, "m2")
    return CATEGORY_PURCHASE, _price_unit_from_basis(basis, "個")


def unit_code_from_sheet(title: str) -> str:
    """シート名 'BFQ7V (2025.9)' / 'BFQ3V_2026.7' / 'PLD11-4R62 (2025.4.1)' → 型式コード"""
    s = title.strip()
    s = re.sub(r"\s*[\(（][^)）]*[\)）]\s*$", "", s)
    s = re.sub(r"[_\s]\d{4}\.\d+(\.\d+)?$", "", s)
    return s.strip()


# ------------------------------------------------------------
# 単価マスタシート
# ------------------------------------------------------------
def _find_master_sheet(wb):
    for ws in wb.worksheets:
        if isinstance(ws["C1"].value, str) and "鋼材単価表" in ws["C1"].value:
            return ws
    return None


def parse_master(ws_v, ws_f) -> dict:
    """単価マスタシートを読む。ws_v=値（data_only）, ws_f=数式。

    戻り値:
      materials: OrderedDict[key] = {name, category, price_unit, steel_grade, thickness_mm, stock_size,
                                     basis, note, prices:[{date, price}], aliases:[...]}
      alias_map: 原価表上名称 → key
      price_dates: {old: date, new: date}
      warnings: [...]
    """
    materials = OrderedDict()
    alias_map = {}
    warnings = []
    old_date = _parse_date_header(ws_v["E2"].value) or dt.date(2020, 2, 1)
    new_date = _parse_date_header(ws_v["F2"].value) or dt.date(2025, 9, 1)

    def add_prices(m, old_p, new_p):
        if _is_num(old_p):
            m["prices"].append({"date": old_date, "price": float(old_p)})
        if _is_num(new_p):
            m["prices"].append({"date": new_date, "price": float(new_p)})

    # --- 上段: 鋼材単価表 ---
    r = 3
    max_r = ws_v.max_row
    while r <= max_r:
        b, c, d = ws_v.cell(r, 2).value, ws_v.cell(r, 3).value, ws_v.cell(r, 4).value
        if isinstance(b, str) and ("部品リスト" in b):
            break
        if isinstance(c, str) and c.strip() and _is_num(d):
            grade = c.strip()
            key = f"STEEL|{grade}|{norm_label(d)}"
            is_area = grade.startswith(("ﾊﾟﾝﾁﾝｸﾞ", "パンチング", "ｴｷｽﾊﾟﾝﾄﾞ"))
            m = materials.get(key)
            if not m:
                m = {
                    "key": key, "name": f"{grade} {norm_label(d)}t",
                    "category": CATEGORY_AREA if is_area else CATEGORY_STEEL,
                    "price_unit": "m2" if is_area else "kg",
                    "steel_grade": grade.split("-")[0], "thickness_mm": float(d),
                    "stock_size": grade.split("-", 1)[1] if "-" in grade else None,
                    "basis": "鋼材単価表", "note": None, "prices": [], "aliases": [],
                }
                materials[key] = m
                add_prices(m, ws_v.cell(r, 5).value, ws_v.cell(r, 6).value)
            # 「引用用①」列（B）の連結キー（例 SPCC-3*61）も別名として登録
            if b is not None:
                alias_map[norm_label(b)] = key
        r += 1

    # --- 下段: 部品リスト ---
    while r <= max_r and ws_v.cell(r, 2).value != "原価表上名称":
        r += 1
    r += 1
    while r <= max_r:
        b = ws_v.cell(r, 2).value
        if b is None or norm_label(b) == "":
            r += 1
            continue
        label = norm_label(b)
        c, d = ws_v.cell(r, 3).value, ws_v.cell(r, 4).value
        old_p, new_p = ws_v.cell(r, 5).value, ws_v.cell(r, 6).value
        basis = ws_v.cell(r, 7).value
        note = ws_v.cell(r, 10).value
        if isinstance(c, str) and c.strip() and _is_num(d):
            key = f"STEEL|{c.strip()}|{norm_label(d)}"
            if key not in materials:
                # 鋼材単価表に無い組合せ（例: SPCE-672*914 1t）は部品リスト側の単価で登録
                grade = c.strip()
                materials[key] = {
                    "key": key, "name": f"{grade} {norm_label(d)}t", "category": CATEGORY_STEEL,
                    "price_unit": "kg", "steel_grade": grade.split("-")[0], "thickness_mm": float(d),
                    "stock_size": grade.split("-", 1)[1] if "-" in grade else None,
                    "basis": basis if isinstance(basis, str) else None, "note": None,
                    "prices": [], "aliases": [],
                }
                add_prices(materials[key], old_p, new_p)
            materials[key]["aliases"].append(label)
            alias_map[label] = key
        else:
            key = "PART|" + label
            cat, unit = classify_material(label, basis if isinstance(basis, str) else "")
            m = materials.get(key)
            if not m:
                m = {
                    "key": key, "name": label, "category": cat, "price_unit": unit,
                    "steel_grade": None, "thickness_mm": None, "stock_size": None,
                    "basis": basis if isinstance(basis, str) else None,
                    "note": note if isinstance(note, str) else None,
                    "prices": [], "aliases": [label],
                }
                materials[key] = m
                add_prices(m, old_p, new_p)
                if not m["prices"]:
                    warnings.append({"kind": "no_price", "sheet": ws_v.title, "row": r, "label": label,
                                     "message": "単価が数値でない（数式エラーか未入力）"})
            alias_map[label] = key
        r += 1

    return {"materials": materials, "alias_map": alias_map,
            "price_dates": {"old": old_date, "new": new_date}, "warnings": warnings}


# ------------------------------------------------------------
# 型式シート
# ------------------------------------------------------------
def _header_columns(ws_v):
    """「内訳」を含むヘッダー行を探し、列の役割を見出し文字から決める"""
    for r in range(1, min(ws_v.max_row, 60) + 1):
        texts = {}
        for c in range(1, ws_v.max_column + 1):
            v = ws_v.cell(r, c).value
            if isinstance(v, str):
                texts[c] = v.replace("　", " ").strip()
        if any(t == "内訳" for t in texts.values()):
            cols = {"name": None, "weight": None, "measure": None, "qty": None,
                    "price": [], "cost": [], "markup": None, "note": None}
            for c, t in texts.items():
                if "板厚" in t and cols["name"] is None:
                    cols["name"] = c
                elif "歩留まり" in t:
                    cols["weight"] = c
                elif t.startswith("使用"):
                    cols["measure"] = c
                elif t.startswith("数量"):
                    cols["qty"] = c
                elif "原価/購入価格" in t or "購入価格" in t:
                    cols["price"].append(c)
                elif t == "原価価格":
                    cols["cost"].append(c)
                elif "倍率" in t:
                    cols["markup"] = c
                elif t == "備考":
                    cols["note"] = c
            if cols["name"] and cols["cost"]:
                return r, cols
    return None, None


def _sheet_refs(formula: str):
    return [m.group(1).strip() for m in _SHEET_REF_RE.finditer(formula)] if isinstance(formula, str) else []


def _find_total_cell(ws_f, ws_v):
    """型式の合計セル。原則 D2（=SUM(D6:D31) / =C6+C7+... / =D6+D7+...）。無ければ先頭 5 行を探す。"""
    from openpyxl.utils import column_index_from_string
    best = None
    for r in range(1, 6):
        for c in range(2, 9):
            f = ws_f.cell(r, c).value
            if not (isinstance(f, str) and f.startswith("=") and "!" not in f):
                continue
            refs = _CELL_RE.findall(f)
            if not ("SUM(" in f.upper() or len(refs) >= 3):
                continue
            cand = (r, c)
            if (r, c) == (2, 4):
                return cand
            if best is None:
                best = cand
    return best


def _base_rows(ws_f, cost_letter: str, total_cell) -> set:
    """合計セルから式をたどり、合計に含まれる原価列の行を集める（オプション部位を区別するため）。"""
    from openpyxl.utils import get_column_letter, column_index_from_string
    rows = set()
    if total_cell is None:
        return rows
    work = [total_cell]
    seen = set()
    while work:
        r, c = work.pop()
        if (r, c) in seen or r < 1 or c < 1:
            continue
        seen.add((r, c))
        f = ws_f.cell(r, c).value
        letter = get_column_letter(c)
        if letter == cost_letter and not (isinstance(f, str) and f.upper().startswith("=SUM(")):
            rows.add(r)
        if not (isinstance(f, str) and f.startswith("=")):
            continue
        if "!" in f:
            continue
        for c1, r1, c2, r2 in _SUM_RE.findall(f):
            if c1.upper() == cost_letter:
                rows.update(range(int(r1), int(r2) + 1))
            else:
                ci = column_index_from_string(c1.upper())
                for rr in range(int(r1), int(r2) + 1):
                    work.append((rr, ci))
        f_wo_sum = _SUM_RE.sub("", f)
        for cl, rw in _CELL_RE.findall(f_wo_sum):
            work.append((int(rw), column_index_from_string(cl.upper())))
    return rows


def parse_unit_sheet(ws_v, ws_f, sheet_names: set) -> dict | None:
    header_row, cols = _header_columns(ws_v)
    if header_row is None:
        return None
    from openpyxl.utils import get_column_letter
    cost_col = cols["cost"][-1]
    old_cost_col = cols["cost"][0] if len(cols["cost"]) > 1 else None
    price_col = cols["price"][-1] if cols["price"] else None
    cost_letter = get_column_letter(cost_col)

    # 小計 SUM の範囲と、原価列を直接参照しているセル（ヘッダーの =G195 等）を集める
    covered, direct = set(), set()
    for r in range(1, ws_f.max_row + 1):
        for c in range(1, ws_f.max_column + 1):
            f = ws_f.cell(r, c).value
            if not (isinstance(f, str) and f.startswith("=")):
                continue
            for c1, r1, c2, r2 in _SUM_RE.findall(f):
                if c1.upper() == cost_letter:
                    covered.update(range(int(r1), int(r2) + 1))
            if "!" not in f:
                for cl, rw in _CELL_RE.findall(f):
                    if cl.upper() == cost_letter:
                        direct.add(int(rw))

    # 合計セルからたどれる行 = 標準構成。SUM には入っているが合計に入らない部位 = オプション
    total_cell = _find_total_cell(ws_f, ws_v)
    base_rows = _base_rows(ws_f, cost_letter, total_cell)

    # 型式名: 先頭数行の英字を含む文字列
    display = None
    for addr in ("A1", "A2", "B2", "B1"):
        v = ws_v[addr].value
        if isinstance(v, str) and re.search(r"[A-Za-z]", v) and len(v) < 40:
            display = v.strip()
            break
    total_v = ws_v.cell(*total_cell).value if total_cell else None
    unit = {
        "sheet": ws_v.title,
        "unit_code": unit_code_from_sheet(ws_v.title),
        "unit_name": display or unit_code_from_sheet(ws_v.title),
        "excel_total": float(total_v) if _is_num(total_v) else None,
        "excel_total_old": ws_v["C2"].value if _is_num(ws_v["C2"].value) else None,
        "lines": [], "warnings": [],
    }
    section = None
    for r in range(header_row + 1, ws_v.max_row + 1):
        name_v = ws_v.cell(r, cols["name"]).value
        cost_f = ws_f.cell(r, cost_col).value
        cost_v = ws_v.cell(r, cost_col).value
        # 部位見出し・TECHS 品番（名称列より左）
        techs = None
        for c in range(1, cols["name"]):
            v = ws_v.cell(r, c).value
            if _is_num(v) and float(v).is_integer() and _TECHS_RE.match(str(int(v))):
                techs = str(int(v))
            elif isinstance(v, str):
                s = v.strip()
                if _TECHS_RE.match(s):
                    techs = s
                elif s and s not in ("techs品番",) and not s.startswith("="):
                    if len(s) > 1 or section is None:
                        section = s
        if name_v is None or norm_label(name_v) == "":
            continue
        label = norm_label(name_v)
        if label in ("内訳", "板厚(mm)/ﾌﾗﾝｼﾞ/購入部品"):
            continue
        w = ws_v.cell(r, cols["weight"]).value if cols["weight"] else None
        w_f = ws_f.cell(r, cols["weight"]).value if cols["weight"] else None
        ms = ws_v.cell(r, cols["measure"]).value if cols["measure"] else None
        ms_f = ws_f.cell(r, cols["measure"]).value if cols["measure"] else None
        q = ws_v.cell(r, cols["qty"]).value if cols["qty"] else None
        q_f = ws_f.cell(r, cols["qty"]).value if cols["qty"] else None
        price_f = ws_f.cell(r, price_col).value if price_col else None
        price_v = ws_v.cell(r, price_col).value if price_col else None

        refs = [s for s in _sheet_refs(price_f) + _sheet_refs(cost_f) if s in sheet_names and s != ws_v.title]
        line = {
            "row": r, "section": section, "label": label, "techs_code": techs,
            "line_type": None, "qty": None, "qty_expr": None, "sub_unit_code": None,
            "excel_price": float(price_v) if _is_num(price_v) else None,
            "excel_cost": float(cost_v) if _is_num(cost_v) else None,
            "cost_error": isinstance(cost_v, str) and cost_v.startswith("#"),
            "markup": float(ws_v.cell(r, cols["markup"]).value) if cols["markup"] and _is_num(ws_v.cell(r, cols["markup"]).value) else 1.0,
            "note": norm_label(ws_v.cell(r, cols["note"]).value) if cols["note"] and ws_v.cell(r, cols["note"]).value is not None else None,
            "orphan": False, "is_option": False,
        }
        if refs:
            line["line_type"] = "subassembly"
            line["sub_unit_code"] = unit_code_from_sheet(refs[0])
            line["qty"] = float(q) if _is_num(q) else 1.0
            line["qty_expr"] = q_f if isinstance(q_f, str) and q_f.startswith("=") else None
        elif _is_num(w):
            line["line_type"] = "weight"
            line["qty"] = float(w)
            line["qty_expr"] = w_f if isinstance(w_f, str) and w_f.startswith("=") else None
            if _is_num(q) and float(q) not in (0.0, 1.0):
                # 1 個あたり重量 × 個数（例: セルプレートリング 2.26kg × 36 個）
                line["qty"] = float(w) * float(q)
                line["qty_expr"] = f"={norm_label(w)}*{norm_label(q)}"
                line["pieces"] = float(q)
        elif _is_num(ms):
            line["line_type"] = "measure"
            line["qty"] = float(ms)
            line["qty_expr"] = ms_f if isinstance(ms_f, str) and ms_f.startswith("=") else None
        elif _is_num(q):
            if "/外注/" in label or label.startswith("外注/"):
                line["line_type"] = "subcontract"
            elif (section and "塗料" in section) or _PAINT_RE.search(label):
                line["line_type"] = "paint"
            else:
                line["line_type"] = "count"
            line["qty"] = float(q)
            line["qty_expr"] = q_f if isinstance(q_f, str) and q_f.startswith("=") else None
        elif isinstance(cost_f, str) and cost_f.startswith("=SUM("):
            continue  # 小計行
        else:
            # 数量の無い行（板厚だけ書いて重量が空の予備行 等）は明細にしない
            if not _is_num(cost_v) or float(cost_v) == 0.0:
                continue
            line["line_type"] = "count"
            line["qty"] = 1.0
        if r in base_rows:
            pass
        elif r in covered or r in direct:
            line["is_option"] = True
        else:
            line["orphan"] = True
            unit["warnings"].append({"kind": "orphan", "sheet": ws_v.title, "row": r, "label": label,
                                     "message": "部位小計の SUM 範囲に含まれていない行（Excel の合計に入っていない）"})
        if line["cost_error"]:
            unit["warnings"].append({"kind": "cost_error", "sheet": ws_v.title, "row": r, "label": label,
                                     "message": f"原価セルがエラー（{cost_v}）。単価マスタに名称が無い可能性"})
        unit["lines"].append(line)
    return unit


# ------------------------------------------------------------
# ブック全体
# ------------------------------------------------------------
def parse_workbook(path_or_file) -> dict:
    """原価ブックを解析して dict を返す（DB 非依存）"""
    if openpyxl is None:
        raise RuntimeError("openpyxl がインストールされていません")
    wb_v = openpyxl.load_workbook(path_or_file, data_only=True)
    if hasattr(path_or_file, "seek"):
        path_or_file.seek(0)
    wb_f = openpyxl.load_workbook(path_or_file, data_only=False)
    ws_m = _find_master_sheet(wb_v)
    if ws_m is None:
        raise ValueError("単価マスタシート（C1 に『鋼材単価表』）が見つかりません")
    master = parse_master(ws_m, wb_f[ws_m.title])
    sheet_codes = {ws.title for ws in wb_v.worksheets}
    units = []
    for ws in wb_v.worksheets:
        if ws.title == ws_m.title:
            continue
        u = parse_unit_sheet(ws, wb_f[ws.title], sheet_codes)
        if u and u["lines"]:
            units.append(u)
    # 明細の名称を単価マスタに突合し、未登録を警告
    unresolved = []
    for u in units:
        for ln in u["lines"]:
            if ln["line_type"] == "subassembly":
                continue
            key = master["alias_map"].get(ln["label"])
            ln["material_key"] = key
            if key is None:
                unresolved.append({"kind": "unresolved", "sheet": u["sheet"], "row": ln["row"], "label": ln["label"],
                                   "message": "単価マスタに名称が無い（表記ゆれの可能性）"})
    return {
        "master": master, "units": units,
        "warnings": master["warnings"] + unresolved + [w for u in units for w in u["warnings"]],
        "unit_codes": [u["unit_code"] for u in units],
    }


def build_context(parsed: dict, include_orphans: bool = False) -> dict:
    """解析結果から costing_engine 用のコンテキスト（メモリ上）を組み立てる。
    取込プレビューで Excel の合計と突合するために使う。id は素材キー／型式コードをそのまま使う。"""
    materials, prices = {}, {}
    for key, m in parsed["master"]["materials"].items():
        materials[key] = {
            "id": key, "code": temp_code(m["name"]), "name": m["name"], "category": m["category"],
            "price_unit": m["price_unit"], "preferred_supplier_id": None, "price_policy": "preferred",
        }
        prices[key] = [{"supplier_id": None, "effective_date": p["date"], "price": p["price"]} for p in m["prices"]]
    units = {}
    for u in parsed["units"]:
        lines = []
        option_lines = {}
        for ln in u["lines"]:
            if ln["orphan"] and not include_orphans:
                continue
            key = ln.get("material_key")
            if ln["line_type"] != "subassembly" and key is None:
                key = "PART|" + ln["label"]
                if key not in materials:
                    cat, unit = classify_material(ln["label"])
                    materials[key] = {"id": key, "code": temp_code(ln["label"]), "name": ln["label"],
                                      "category": cat, "price_unit": unit, "preferred_supplier_id": None,
                                      "price_policy": "preferred"}
                    prices[key] = []
            line = {
                "id": f"{u['unit_code']}#{ln['row']}", "section": ln["section"], "label": ln["label"],
                "line_type": ln["line_type"], "material_id": None if ln["line_type"] == "subassembly" else key,
                "sub_unit_id": ln["sub_unit_code"] if ln["line_type"] == "subassembly" else None,
                "qty": ln["qty"], "qty_expr": ln["qty_expr"], "markup_factor": ln["markup"], "note": ln["note"],
            }
            if ln.get("is_option"):
                # 合計に入らない部位（BFQ の「各オプション」）は別ユニットにする
                option_lines.setdefault(ln["section"] or "オプション", []).append(line)
            else:
                lines.append(line)
        units[u["unit_code"]] = {"id": u["unit_code"], "code": u["unit_code"], "name": u["unit_name"], "lines": lines}
        for sec, ols in option_lines.items():
            code = f"{u['unit_code']}/{sec}"
            units[code] = {"id": code, "code": code, "name": f"{u['unit_name']} {sec}（オプション）", "lines": ols,
                           "is_option": True, "parent_code": u["unit_code"]}
    return {"materials": materials, "prices": prices, "adjustments": {}, "units": units}
